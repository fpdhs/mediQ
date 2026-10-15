import os
import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\user\Downloads\Medi Q 중간발표.pptx"
prs = pptx.Presentation(pptx_path)
print("Total slides:", len(prs.slides))
print(f"Slide dimensions: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} inches")

for idx, slide in enumerate(prs.slides):
    texts = []
    has_image = False
    has_media = False
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    texts.append(t)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            has_image = True
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.MEDIA:
            has_media = True

    short_texts = [t for t in texts if len(t) < 60]
    title = short_texts[0] if short_texts else "(No title)"
    print(f"Slide {idx + 1:02d}: {title}")
    if len(short_texts) > 1:
        print(f"          Sub: {short_texts[1:4]}")
