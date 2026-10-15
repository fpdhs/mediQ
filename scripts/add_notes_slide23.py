import pptx

pptx_path = r"C:\Users\user\Downloads\Medi Q 중간발표.pptx"
prs = pptx.Presentation(pptx_path)
slide23 = prs.slides[22]

notes_text = """[MediQ 아키텍처 모션 비디오 발표 스크립트]
"지금 보시는 영상은 MediQ의 핵심 아키텍처 전 주기 흐름을 시각화한 모션 데모입니다.

❶ [신원 인증]: 환자와 의료진이 Microsoft Entra ID를 통해 신원을 인증합니다. 인증(Who)과 의료영상 접근 인가(What)가 물리적으로 분리되어 있어 계정이 탈취되더라도 영상이 유출되지 않습니다.
❷ [환자 동의]: 환자가 모바일에서 본인의 서울성모병원 복부 CT 700장을 확인하고, 수신 병원(아산병원)과 열람 전용(VIEW_ONLY) 권한, 7일의 유효기간을 직접 지정합니다.
❸ [QR 접수증]: 동의가 체결되면 위변조가 불가능한 HMAC-SHA256 암호화 일회용 QR 티켓이 발급됩니다.
❹ [현장 접수]: 수신병원 진료실에서 의료진이 환자가 제시한 QR 코드를 원터치 스캔하여 조회를 요청합니다.
❺ [ABAC 11단계 검증]: 중앙 Control Plane에서 11개 교집합 조건(동의 활성, 병원 일치, 의사 소속, 허용 Study UID 등)을 전수 검증하며, 단 하나라도 불일치 시 Fail-Closed로 즉시 차단합니다.
❻ [단기 토큰]: 11단계 검증을 모두 통과하면 5분 유효 단기 토큰이 발급됩니다.
❼ [게이트웨이 재검증]: B병원 분리 게이트웨이에서 DPoP 및 토큰 서명을 재확인합니다.
❽ [700장 CT 직통 스트리밍]: 클라우드 중앙에 원본을 저장하지 않는 'Zero Cloud Storage' 원칙에 따라, WireGuard 사설망을 통해 B병원 뷰어로 700장 CT가 무지연 스트리밍 렌더링됩니다."
"""

notes_slide = slide23.notes_slide
text_frame = notes_slide.notes_text_frame
text_frame.text = notes_text

prs.save(pptx_path)
print("Successfully injected speaker notes into Slide 23!")
