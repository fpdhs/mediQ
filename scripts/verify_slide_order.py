import pptx
import sys

sys.stdout.reconfigure(encoding='utf-8')

pptx_path = r"C:\Users\user\Downloads\Medi Q 중간발표.pptx"
prs = pptx.Presentation(pptx_path)

print(f"Total slides count: {len(prs.slides)}")

for idx in [21, 22, 23, 24]:
    slide = prs.slides[idx]
    texts = []
    media_name = ""
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t and len(t) < 60:
                    texts.append(t)
        if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.MEDIA:
            media_name = shape.name

    title = texts[0] if texts else (f"[MEDIA: {media_name}]" if media_name else "(No text)")
    notes = slide.notes_slide.notes_text_frame.text[:80].replace('\n', ' ') if slide.has_notes_slide else "(No notes)"
    print(f"Slide {idx + 1:02d}: Title = {title}")
    if media_name:
        print(f"          Media = {media_name}")
    print(f"          Notes = {notes}")
