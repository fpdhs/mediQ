import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

PPTX_PATH = r"C:\Users\user\Downloads\Medi Q 중간발표.pptx"
QR_PATH = r"C:\Users\user\Documents\New project\public\demo-assets\mediq_live_qr.png"

FONT = "프리텐다드"

# MediQ Palette
C_NAVY     = RGBColor(11, 35, 66)    # #0B2342
C_BLUE     = RGBColor(47, 104, 181)  # #2F68B5
C_BLUE2    = RGBColor(77, 134, 207)  # #4D86CF
C_CYAN     = RGBColor(85, 194, 216)  # #55C2D8
C_TEAL     = RGBColor(22, 139, 130)  # #168B82
C_INK      = RGBColor(23, 38, 58)    # #17263A
C_SLATE    = RGBColor(83, 101, 123)  # #53657B
C_MUTED    = RGBColor(123, 138, 155) # #7B8A9B
C_LINE     = RGBColor(216, 224, 233) # #D8E0E9
C_PALE     = RGBColor(243, 246, 250) # #F3F6FA
C_PALEBLUE = RGBColor(233, 241, 252) # #E9F1FC
C_PALETEAL = RGBColor(231, 244, 241) # #E7F4F1
C_WHITE    = RGBColor(255, 255, 255) # #FFFFFF

def px(val):
    return Inches(val / 96.0)

def pt(px_val):
    return Pt(px_val * 0.75)

def main():
    if not os.path.exists(PPTX_PATH):
        raise FileNotFoundError(f"PPTX not found: {PPTX_PATH}")
    if not os.path.exists(QR_PATH):
        raise FileNotFoundError(f"QR code image not found: {QR_PATH}")

    prs = pptx.Presentation(PPTX_PATH)
    print(f"Original presentation loaded: {len(prs.slides)} slides.")

    # 1. Add new blank slide (appends to end)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # 2. Build Slide Background & Header
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(1920), px(1080))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_WHITE
    bg.line.fill.background()

    hdr_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, px(1920), px(230))
    hdr_bg.fill.solid()
    hdr_bg.fill.fore_color.rgb = C_PALE
    hdr_bg.line.fill.background()

    # Blue vertical accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(80), px(50), px(8), px(130))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_BLUE
    bar.line.fill.background()

    # Section Label
    lbl_box = slide.shapes.add_textbox(px(110), px(44), px(1400), px(34))
    tf = lbl_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "04  |  SYSTEM DESIGN"
    p.font.name = FONT
    p.font.size = pt(20)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    # Title
    t_box = slide.shapes.add_textbox(px(110), px(80), px(1400), px(60))
    tf = t_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "실시간 인터랙티브 시연 & 라이브 서버 접속"
    p.font.name = FONT
    p.font.size = pt(36)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Subtitle
    sub_box = slide.shapes.add_textbox(px(110), px(145), px(1650), px(40))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "동일 네트워크(Wi-Fi / 사내망) 실시간 접속 환경  |  QR 코드 원터치 모바일 스캔 및 브라우저 즉시 체험"
    p.font.name = FONT
    p.font.size = pt(18)
    p.font.color.rgb = C_SLATE

    # Top Right Live Badge in Header
    hdr_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(1470), px(65), px(370), px(48))
    hdr_pill.fill.solid()
    hdr_pill.fill.fore_color.rgb = C_PALETEAL
    hdr_pill.line.color.rgb = C_TEAL
    hdr_pill.line.width = Pt(1.5)
    tf = hdr_pill.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "🟢 LIVE ACTIVE (PORT 3098)"
    p.font.name = FONT
    p.font.size = pt(15)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # 3. Left Panel: Live Access & QR Code
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(80), px(255), px(760), px(760))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = C_WHITE
    left_card.line.color.rgb = C_LINE
    left_card.line.width = Pt(1.2)

    # Server Status Pill
    status_pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(120), px(285), px(680), px(50))
    status_pill.fill.solid()
    status_pill.fill.fore_color.rgb = C_PALE
    status_pill.line.color.rgb = C_BLUE2
    status_pill.line.width = Pt(1.0)
    tf = status_pill.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "사내망 / 동일 Wi-Fi 접속 대역: 10.10.70.x  (0.0.0.0 Listen)"
    p.font.name = FONT
    p.font.size = pt(15)
    p.font.bold = True
    p.font.color.rgb = C_INK

    # QR Code Frame
    qr_frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(245), px(360), px(430), px(430))
    qr_frame.fill.solid()
    qr_frame.fill.fore_color.rgb = C_PALEBLUE
    qr_frame.line.color.rgb = C_LINE
    qr_frame.line.width = Pt(1.0)

    # QR Code Image
    slide.shapes.add_picture(QR_PATH, px(265), px(380), px(390), px(390))

    # QR Guidance Text
    qr_txt = slide.shapes.add_textbox(px(120), px(805), px(680), px(35))
    tf = qr_txt.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "📱 스마트폰 기본 카메라로 QR 코드를 스캔하세요"
    p.font.name = FONT
    p.font.size = pt(18)
    p.font.bold = True
    p.font.color.rgb = C_INK

    # Main Clickable URL Badge Button
    url_btn = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(120), px(850), px(680), px(78))
    url_btn.fill.solid()
    url_btn.fill.fore_color.rgb = C_PALEBLUE
    url_btn.line.color.rgb = C_BLUE
    url_btn.line.width = Pt(2.0)
    tf = url_btn.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "http://10.10.70.151:3098"
    p.font.name = FONT
    p.font.size = pt(25)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    # Add hyperlink to URL button
    url_btn.click_action.hyperlink.address = "http://10.10.70.151:3098"

    # Local fallback label
    local_txt = slide.shapes.add_textbox(px(120), px(942), px(680), px(30))
    tf = local_txt.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "* 발표 PC 직접 접속: http://localhost:3098/architecture-motion.html"
    p.font.name = FONT
    p.font.size = pt(14)
    p.font.color.rgb = C_MUTED

    # 4. Right Top Card: Feature 01 (Architecture Motion Simulator)
    r1_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(870), px(255), px(970), px(365))
    r1_card.fill.solid()
    r1_card.fill.fore_color.rgb = C_PALEBLUE
    r1_card.line.color.rgb = C_LINE
    r1_card.line.width = Pt(1.2)

    # Feature 01 Badge
    f1_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(905), px(285), px(380), px(36))
    f1_badge.fill.solid()
    f1_badge.fill.fore_color.rgb = C_BLUE
    f1_badge.line.fill.background()
    tf = f1_badge.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "FEATURE 01 · 실시간 아키텍처 시뮬레이터"
    p.font.name = FONT
    p.font.size = pt(13)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Feature 01 Path
    f1_path = slide.shapes.add_textbox(px(1300), px(287), px(510), px(32))
    tf = f1_path.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "/architecture-motion.html"
    p.font.name = FONT
    p.font.size = pt(14)
    p.font.bold = True
    p.font.color.rgb = C_BLUE2

    # Feature 01 Title
    f1_title = slide.shapes.add_textbox(px(905), px(330), px(900), px(45))
    tf = f1_title.text_frame
    p = tf.paragraphs[0]
    p.text = "8단계 데이터 흐름 & 통제 원칙 동적 시각화"
    p.font.name = FONT
    p.font.size = pt(22)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Feature 01 Bullets
    f1_box = slide.shapes.add_textbox(px(905), px(380), px(900), px(215))
    tf = f1_box.text_frame
    tf.word_wrap = True

    bullets1 = [
        ("✔  11단계 ABAC 인가 검증", " : 동의 유효성·의료진 소속·Study UID 전수 검증 및 Fail-Closed 즉시 차단"),
        ("✔  Zero Cloud Storage Invariant", " : 클라우드 중앙 무저장 원칙 및 WireGuard 사설망 직통 WADO-RS"),
        ("✔  인터랙티브 카메라 줌·팬", " : 8개 시스템 컴포넌트 실시간 포커스, 꼬리 잔상 코멧 파티클 및 1080p 녹화")
    ]
    for idx, (b_title, b_desc) in enumerate(bullets1):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = pt(12)
        run1 = p.add_run()
        run1.text = b_title
        run1.font.name = FONT
        run1.font.size = pt(16)
        run1.font.bold = True
        run1.font.color.rgb = C_BLUE

        run2 = p.add_run()
        run2.text = b_desc
        run2.font.name = FONT
        run2.font.size = pt(15)
        run2.font.color.rgb = C_INK

    # 5. Right Bottom Card: Feature 02 (3-Role Demo Studio)
    r2_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(870), px(650), px(970), px(365))
    r2_card.fill.solid()
    r2_card.fill.fore_color.rgb = C_PALETEAL
    r2_card.line.color.rgb = C_LINE
    r2_card.line.width = Pt(1.2)

    # Feature 02 Badge
    f2_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(905), px(680), px(360), px(36))
    f2_badge.fill.solid()
    f2_badge.fill.fore_color.rgb = C_TEAL
    f2_badge.line.fill.background()
    tf = f2_badge.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "FEATURE 02 · 3대 역할 전 주기 시연"
    p.font.name = FONT
    p.font.size = pt(13)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    # Feature 02 Path
    f2_path = slide.shapes.add_textbox(px(1300), px(682), px(510), px(32))
    tf = f2_path.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    p.text = "/demo-showcase.html"
    p.font.name = FONT
    p.font.size = pt(14)
    p.font.bold = True
    p.font.color.rgb = C_TEAL

    # Feature 02 Title
    f2_title = slide.shapes.add_textbox(px(905), px(725), px(900), px(45))
    tf = f2_title.text_frame
    p = tf.paragraphs[0]
    p.text = "환자 · 의료진 · 보안관리자 엔드투엔드 시연"
    p.font.name = FONT
    p.font.size = pt(22)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Feature 02 Bullets
    f2_box = slide.shapes.add_textbox(px(905), px(775), px(900), px(215))
    tf = f2_box.text_frame
    tf.word_wrap = True

    bullets2 = [
        ("👤  환자 포털", " : Entra ID 본인인증 ➔ 복부 CT 700장 선택 ➔ VIEW_ONLY 동의 ➔ QR 티켓 발급"),
        ("🩺  의료진 뷰어", " : QR 스캔 접수 ➔ 5분 단기 토큰 발급 ➔ 로컬 PACS 직통 700장 CT 무지연 스트리밍"),
        ("🛡️  보안관리자", " : 전 주기 SHA-256 해시 체인 불변 감사로그 & 대량 반출 실시간 차단 모니터링")
    ]
    for idx, (b_title, b_desc) in enumerate(bullets2):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = pt(12)
        run1 = p.add_run()
        run1.text = b_title
        run1.font.name = FONT
        run1.font.size = pt(16)
        run1.font.bold = True
        run1.font.color.rgb = C_TEAL

        run2 = p.add_run()
        run2.text = b_desc
        run2.font.name = FONT
        run2.font.size = pt(15)
        run2.font.color.rgb = C_INK

    # 6. Inject Speaker Notes for Slide 23
    notes_text = """[Slide 23 발표 스크립트: 실시간 라이브 서버 접속 안내]
"다음으로, 저희가 구축한 시스템의 실시간 동작을 직접 확인하실 수 있도록 라이브 서버를 시연해 드리겠습니다.
지금 발표장에 계신 심사위원님들과 청중 여러분께서도 스마트폰 기본 카메라로 화면 좌측의 QR 코드를 스캔하시면, 별도의 앱 설치 없이 동일 Wi-Fi 망(http://10.10.70.151:3098)을 통해 MediQ 실시간 아키텍처 시뮬레이터에 즉시 접속하실 수 있습니다.
이 시뮬레이터에서는 환자 동의 생성부터 11단계 ABAC 인가, 그리고 병원 간 WireGuard 사설 터널을 통한 700장 CT 스트리밍까지 전 과정을 인터랙티브하게 체험하실 수 있습니다.
그럼 다음 슬라이드에서 실제 아키텍처가 어떻게 흐르는지 고화질 영상으로 확인하시겠습니다."
"""
    slide.notes_slide.notes_text_frame.text = notes_text

    # 7. Move this new slide to index 22 (so it becomes Slide 23)
    sldIdLst = prs.slides._sldIdLst
    new_sldId = sldIdLst[-1]
    sldIdLst.remove(new_sldId)
    sldIdLst.insert(22, new_sldId)

    # 8. Save
    prs.save(PPTX_PATH)
    print("Slide 23 inserted successfully at position 23!")
    print(f"Total slides now: {len(prs.slides)}")

if __name__ == '__main__':
    main()
